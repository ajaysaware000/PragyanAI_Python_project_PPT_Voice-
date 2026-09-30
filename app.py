import streamlit as st
from pptx import Presentation
from gtts import gTTS
import io

# Page configuration
st.set_page_config(
    page_title="PPT to Voice Converter",
    page_icon="🔊",
    layout="centered"
)

st.title("📊 PPT to Voice Converter")
st.write("Upload a PowerPoint, select a slide, and convert its text to MP3.")

# Upload PPT
uploaded_file = st.file_uploader(
    "Upload your PowerPoint file",
    type=["pptx"]
)

if uploaded_file is not None:

    # Read presentation
    presentation = Presentation(uploaded_file)

    total_slides = len(presentation.slides)

    st.success(f"✅ PPT uploaded successfully!")
    st.write(f"**Total slides:** {total_slides}")

    # Select slide
    slide_number = st.number_input(
        "Select slide number",
        min_value=1,
        max_value=total_slides,
        value=1,
        step=1
    )

    # Extract selected slide
    slide = presentation.slides[slide_number - 1]

    slide_text_parts = []

    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text = shape.text.strip()

            if text:
                slide_text_parts.append(text)

    slide_text = "\n".join(slide_text_parts)

    st.subheader(f"📝 Slide {slide_number} Text")

    if slide_text:

        st.text_area(
            "Extracted text",
            slide_text,
            height=250
        )

        # Convert to voice
        if st.button("🔊 Convert to MP3"):

            with st.spinner("Converting text to voice..."):

                try:
                    tts = gTTS(
                        text=slide_text,
                        lang="en",
                        slow=False
                    )

                    audio_buffer = io.BytesIO()
                    tts.write_to_fp(audio_buffer)

                    audio_buffer.seek(0)

                    st.success("✅ Voice generated successfully!")

                    # Play audio
                    st.audio(
                        audio_buffer,
                        format="audio/mp3"
                    )

                    # Download button
                    st.download_button(
                        label="⬇️ Download MP3",
                        data=audio_buffer,
                        file_name=f"slide_{slide_number}.mp3",
                        mime="audio/mpeg"
                    )

                except Exception as e:

                    st.error(
                        f"❌ Error while converting text to voice: {e}"
                    )

    else:

        st.warning(
            "⚠️ No text was found on this slide."
        )
